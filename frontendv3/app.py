"""Akshar Setu — Gujarati OCR using the local Bodhan AI Indic-OCR pipeline."""
from __future__ import annotations

import json
import os
import re
import sys
import tempfile
import threading
import time
from pathlib import Path
from typing import Any, Iterator

import gradio as gr
import pymupdf
from PIL import Image, ImageOps
from docx import Document


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------
MODEL_PATH = os.getenv(
    "INDIC_OCR_MODEL_PATH",
    r"D:\code material 2.0\akshar-setu-v3\models\indic-ocr",
)
PDF_RENDER_DPI = 200
MAX_PDF_PAGES = 300
MAX_IMAGE_PIXELS = 50_000_000
SUPPORTED_IMAGES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
ASSETS_DIR = Path(__file__).resolve().parent / "assets"
FAVICON_PATH = ASSETS_DIR / "favicon.png"

_MODEL = None
_MODEL_LOCK = threading.Lock()


# ---------------------------------------------------------------------------
# Existing local Indic-OCR integration — no Tesseract fallback
# ---------------------------------------------------------------------------
def load_ocr_model():
    global _MODEL
    with _MODEL_LOCK:
        if _MODEL is not None:
            return _MODEL

        model_dir = Path(MODEL_PATH).resolve()
        if not model_dir.is_dir():
            raise FileNotFoundError(
                f"Indic-OCR model folder does not exist:\n{model_dir}\n\n"
                "Update MODEL_PATH in app.py or set INDIC_OCR_MODEL_PATH."
            )

        model_dir_str = str(model_dir)
        if model_dir_str not in sys.path:
            sys.path.insert(0, model_dir_str)

        try:
            from indic_ocr import IndicOCR
        except Exception as exc:
            raise RuntimeError(
                "Could not import IndicOCR from your downloaded repository.\n"
                f"Model folder: {model_dir}\n"
                "Check that the folder contains the indic_ocr package and its dependencies.\n\n"
                f"Original error: {exc}"
            ) from exc

        try:
            _MODEL = IndicOCR.from_pretrained(model_dir_str)
        except Exception as exc:
            _MODEL = None
            raise RuntimeError(
                "Indic-OCR model initialization failed.\n"
                f"Model folder: {model_dir}\n\nOriginal error: {exc}"
            ) from exc
        return _MODEL


def normalize_text(text: str) -> str:
    """Normalize whitespace without stripping Gujarati Unicode characters."""
    if not text:
        return ""
    text = text.replace("▁", " ").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def predict_page(image_path: str) -> str:
    model = load_ocr_model()
    result = model.parse(str(image_path))
    if not isinstance(result, dict):
        raise TypeError(
            "IndicOCR.parse() returned an unexpected result. "
            f"Expected dict, received {type(result).__name__}."
        )
    return normalize_text(result.get("markdown", ""))


def parse_page_range(start_page: int, end_page: int, page_count: int) -> list[int]:
    """Validate inclusive, 1-based page bounds and return zero-based indexes."""
    try:
        start = int(start_page)
        end = int(end_page)
    except (TypeError, ValueError) as exc:
        raise ValueError("Enter valid whole-number start and end page values.") from exc
    if start < 1 or end < 1:
        raise ValueError("Page numbers must be 1 or greater.")
    if start > end:
        raise ValueError("Starting page cannot be greater than ending page.")
    if end > page_count:
        raise ValueError(f"Ending page is {end}, but this file has only {page_count} page(s).")
    return list(range(start - 1, end))


def get_uploaded_path(item: Any) -> str:
    if isinstance(item, str):
        return item
    if hasattr(item, "name"):
        return str(item.name)
    raise ValueError("Could not read the uploaded file. Please upload it again.")


def inspect_upload(file_item):
    """Return a useful file summary and set the default ending page."""
    if not file_item:
        return "Upload one PDF or image to get started.", 1
    try:
        path = Path(get_uploaded_path(file_item))
        if not path.is_file():
            raise FileNotFoundError("The uploaded file is no longer available. Upload it again.")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            with pymupdf.open(str(path)) as pdf:
                count = pdf.page_count
            if count < 1:
                raise ValueError("This PDF contains no pages.")
            if count > MAX_PDF_PAGES:
                raise ValueError(f"This PDF has {count} pages. The configured maximum is {MAX_PDF_PAGES}.")
            return f"**{path.name}** · {count} page(s) detected. Choose the page range to extract.", count
        if suffix in SUPPORTED_IMAGES:
            with Image.open(path) as image:
                width, height = image.size
            if width * height > MAX_IMAGE_PIXELS:
                raise ValueError(f"Image is too large ({width:,} × {height:,} pixels).")
            return f"**{path.name}** · Image detected ({width:,} × {height:,} px). Page range is not needed.", 1
        raise ValueError("Unsupported file type. Choose a PDF, PNG, JPG, TIFF, BMP, or WEBP image.")
    except Exception as exc:
        return f"⚠️ {exc}", 1


def _temporary_png_from_pil(image: Image.Image) -> str:
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as temp_file:
            temp_path = temp_file.name
        image.save(temp_path, format="PNG")
        return temp_path
    except Exception:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)
        raise


def iter_file_pages(file_path: str, start_page: int, end_page: int) -> Iterator[tuple[int, int, str]]:
    """Yield (page number, selected-page count, OCR text) one page at a time."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"Uploaded file is no longer available: {path.name}")
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        with pymupdf.open(str(path)) as pdf:
            page_count = pdf.page_count
            if page_count < 1:
                raise ValueError(f"{path.name} contains no pages.")
            if page_count > MAX_PDF_PAGES:
                raise ValueError(f"{path.name} has {page_count} pages; maximum allowed is {MAX_PDF_PAGES}.")
            selected = parse_page_range(start_page, end_page, page_count)
            scale = PDF_RENDER_DPI / 72.0
            for page_index in selected:
                page = pdf.load_page(page_index)
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(scale, scale),
                    colorspace=pymupdf.csRGB,
                    alpha=False,
                )
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                temp_path = None
                try:
                    temp_path = _temporary_png_from_pil(image)
                    text = predict_page(temp_path)
                finally:
                    image.close()
                    if temp_path and os.path.exists(temp_path):
                        try:
                            os.remove(temp_path)
                        except OSError:
                            pass
                yield page_index + 1, len(selected), text
        return

    if suffix in SUPPORTED_IMAGES:
        with Image.open(path) as source:
            width, height = source.size
            if width * height > MAX_IMAGE_PIXELS:
                raise ValueError(f"Image is too large: {width:,} × {height:,} pixels.")
            image = ImageOps.exif_transpose(source).convert("RGB")
        temp_path = None
        try:
            temp_path = _temporary_png_from_pil(image)
            text = predict_page(temp_path)
        finally:
            image.close()
            if temp_path and os.path.exists(temp_path):
                try:
                    os.remove(temp_path)
                except OSError:
                    pass
        yield 1, 1, text
        return

    raise ValueError("Unsupported file type. Upload a PDF, PNG, JPG, JPEG, TIFF, BMP, or WEBP image.")


def format_elapsed(seconds: float) -> str:
    seconds = max(0, int(seconds))
    hours, remainder = divmod(seconds, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def check_model():
    try:
        load_ocr_model()
        return f"✅ **Indic-OCR is ready.**\n\nModel folder: `{Path(MODEL_PATH).resolve()}`"
    except Exception as exc:
        return f"❌ **Model check failed**\n\n{exc}"


def make_download_files(path: Path, page_results: list[dict[str, Any]], combined_text: str) -> tuple[str, str, str]:
    """Create TXT, DOCX, and structured JSON downloads for the completed OCR run."""
    safe_stem = re.sub(r"[^A-Za-z0-9_-]+", "_", path.stem).strip("_") or "akshar_setu_ocr"

    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8-sig", newline="\n", suffix=".txt",
        prefix=f"{safe_stem}_", delete=False,
    ) as txt_file:
        txt_file.write(combined_text)
        txt_file.write("\n")
        txt_path = txt_file.name

    doc = Document()
    doc.add_heading("Akshar Setu — Gujarati OCR", level=0)
    doc.add_paragraph(f"Source file: {path.name}")
    doc.add_paragraph(f"Pages processed: {len(page_results)}")
    for index, item in enumerate(page_results):
        if index:
            doc.add_page_break()
        doc.add_heading(f"Page {item['page']}", level=1)
        doc.add_paragraph(item["text"] or "[No text detected.]")
    with tempfile.NamedTemporaryFile(suffix=".docx", prefix=f"{safe_stem}_", delete=False) as doc_file:
        docx_path = doc_file.name
    doc.save(docx_path)

    json_payload = {
        "application": "Akshar Setu",
        "source_file": path.name,
        "total_pages_processed": len(page_results),
        "total_characters": sum(item["characters"] for item in page_results),
        "extracted_text": combined_text,
        "pages": page_results,
    }
    with tempfile.NamedTemporaryFile(
        mode="w", encoding="utf-8", newline="\n", suffix=".json",
        prefix=f"{safe_stem}_", delete=False,
    ) as json_file:
        json.dump(json_payload, json_file, ensure_ascii=False, indent=2)
        json_file.write("\n")
        json_path = json_file.name

    return txt_path, docx_path, json_path


def run_ocr(file_item, start_page, end_page, progress=gr.Progress()):
    """Generator callback: yield UI updates after every processed page."""
    if not file_item:
        raise gr.Error("Please upload one PDF or image first.")

    try:
        file_path = get_uploaded_path(file_item)
        path = Path(file_path)
        if not path.is_file():
            raise FileNotFoundError("The uploaded file is no longer available. Upload it again.")
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            with pymupdf.open(str(path)) as pdf:
                page_count = pdf.page_count
            selected_pages = parse_page_range(start_page, end_page, page_count)
            total_pages = len(selected_pages)
        elif suffix in SUPPORTED_IMAGES:
            total_pages = 1
            start_page, end_page = 1, 1
        else:
            raise ValueError("Unsupported file type. Upload a PDF or supported image.")
    except Exception as exc:
        raise gr.Error(str(exc)) from exc

    started = time.perf_counter()
    page_results: list[dict[str, Any]] = []
    status = f"Preparing **{path.name}**…\n\n0 of {total_pages} pages completed · elapsed 00:00:00"
    yield "", [], None, None, None, status

    try:
        for page_number, _, page_text in iter_file_pages(file_path, start_page, end_page):
            page_results.append({
                "file": path.name,
                "page": page_number,
                "characters": len(page_text),
                "text": page_text,
            })
            completed = len(page_results)
            elapsed = format_elapsed(time.perf_counter() - started)
            sections = []
            for result in page_results:
                label = f"===== {result['file']} — Page {result['page']} ====="
                body = result["text"] if result["text"] else "[No text detected.]"
                sections.append(f"{label}\n\n{body}")
            combined_text = "\n\n".join(sections)
            status = (
                f"**Processing:** `{path.name}`\n\n"
                f"**Progress:** {completed} / {total_pages} pages ({completed / total_pages:.0%})\n\n"
                f"**Current page:** {page_number}\n\n"
                f"**Elapsed time:** {elapsed}"
            )
            progress(completed / total_pages, desc=f"Page {page_number} of selected pages")
            yield combined_text, page_results, None, None, None, status

        elapsed = format_elapsed(time.perf_counter() - started)
        combined_text = "\n\n".join(
            f"===== {item['file']} — Page {item['page']} =====\n\n"
            f"{item['text'] if item['text'] else '[No text detected.]'}"
            for item in page_results
        )
        txt_path, docx_path, json_path = make_download_files(path, page_results, combined_text)
        status = (
            "✅ **Extraction complete**\n\n"
            f"**Pages processed:** {len(page_results)} / {total_pages}\n\n"
            f"**Characters extracted:** {sum(item['characters'] for item in page_results):,}\n\n"
            f"**Total time:** {elapsed}\n\n"
            "Downloads are ready: TXT, Word (.docx), and page-by-page JSON."
        )
        progress(1.0, desc="OCR complete")
        yield combined_text, page_results, txt_path, docx_path, json_path, status
    except Exception as exc:
        elapsed = format_elapsed(time.perf_counter() - started)
        error_status = f"❌ **OCR failed**\n\n{exc}\n\nElapsed time: {elapsed}"
        if page_results:
            partial_text = "\n\n".join(
                f"===== {item['file']} — Page {item['page']} =====\n\n"
                f"{item['text'] if item['text'] else '[No text detected.]'}"
                for item in page_results
            )
            yield partial_text, page_results, None, None, None, error_status
        else:
            yield "", [], None, None, None, error_status


# ---------------------------------------------------------------------------
# Frontend
# ---------------------------------------------------------------------------
CSS = r"""
.gradio-container { max-width: 1240px !important; margin: 0 auto !important; padding-bottom: 2rem !important; }
#hero { padding: 1.5rem 1.7rem; border: 1px solid var(--border-color-primary); border-radius: 20px; margin-bottom: .8rem; background: var(--background-fill-secondary); }
#hero h1 { margin: 0 0 .35rem 0; letter-spacing: -.035em; font-size: clamp(1.8rem, 3vw, 2.4rem); }
#hero p { margin: 0; color: var(--body-text-color-subdued); }
#subtle-note { color: var(--body-text-color-subdued); font-size: .92rem; }
#status-panel { border: 1px solid var(--border-color-primary); border-radius: 14px; padding: .8rem 1rem; }
#output-text textarea { height: 520px !important; min-height: 520px !important; max-height: 65vh !important; overflow-y: scroll !important; overflow-x: auto !important; resize: vertical !important; white-space: pre-wrap !important; overflow-wrap: anywhere !important; font-size: 1rem !important; line-height: 1.85 !important; }
#download-stack { gap: .65rem !important; }
#download-stack button { min-height: 48px !important; border-radius: 12px !important; font-weight: 600 !important; }
#output-card { border: 1px solid var(--border-color-primary); border-radius: 16px; padding: .8rem; }
footer { display: none !important; }
@media (max-width: 800px) {
  #output-text textarea { height: 420px !important; min-height: 420px !important; max-height: 60vh !important; }
}
@media (prefers-color-scheme: dark) {
  .gradio-container { --background-fill-primary: #101713 !important; --background-fill-secondary: #17221b !important; --block-background-fill: #17221b !important; --body-text-color: #edf5ef !important; --body-text-color-subdued: #a6b8aa !important; --border-color-primary: #34473a !important; background: #101713 !important; color: #edf5ef !important; }
  #hero { background: linear-gradient(135deg, #172a20, #17221b); border-color: #34473a; }
  input, textarea, select { background: #101713 !important; color: #edf5ef !important; }
}
@media (prefers-color-scheme: light) {
  .gradio-container { --background-fill-primary: #ffffff !important; --background-fill-secondary: #f4faf5 !important; --block-background-fill: #ffffff !important; --body-text-color: #17251b !important; --body-text-color-subdued: #65756a !important; --border-color-primary: #d9e6dc !important; }
  #hero { background: linear-gradient(135deg, #f1faf3, #ffffff); }
}
"""

with gr.Blocks(title="Akshar Setu — Gujarati OCR") as demo:
    gr.Markdown(
        """
        <div id="hero">
          <h1>Akshar Setu</h1>
          <p>Gujarati document digitization · Fine-tuned Indic-OCR</p>
        </div>
        """
    )
    gr.Markdown(
        "Upload one document, select the pages, and extract Gujarati text. "
        "Fine-tuned Indic-OCR model is used.",
        elem_id="subtle-note",
    )

    with gr.Row(equal_height=False):
        with gr.Column(scale=5, min_width=290):
            gr.Markdown("### 1 · Input")
            upload = gr.File(
                label="Choose a PDF or image",
                file_count="single",
                file_types=[".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"],
                type="filepath",
            )
            file_info = gr.Markdown("Upload one PDF or image to get started.")
            with gr.Row():
                start_page = gr.Number(label="Starting page", value=1, minimum=1, precision=0, step=1)
                end_page = gr.Number(label="Ending page", value=1, minimum=1, precision=0, step=1)
            gr.Markdown("Page numbers start at 1. For images, page selection is ignored.", elem_id="subtle-note")
            with gr.Row():
                run_button = gr.Button("Extract Gujarati text", variant="primary", scale=3)
                check_button = gr.Button("Check model", scale=1)
            gr.Markdown("### 2 · Processing")
            status = gr.Markdown("Ready when you are. Upload a file to begin.", elem_id="status-panel")

        with gr.Column(scale=8, min_width=360, elem_id="output-card"):
            gr.Markdown("### 3 · Extracted Gujarati text")
            with gr.Row(equal_height=False):
                with gr.Column(scale=7, min_width=250):
                    output_text = gr.Textbox(
                        label="OCR result",
                        placeholder="Your extracted Gujarati text will appear here…",
                        lines=18,
                        max_lines=18,
                        interactive=False,
                        elem_id="output-text",
                        autoscroll=False,
                    )
                with gr.Column(scale=2, min_width=135, elem_id="download-stack"):
                    gr.Markdown("**Download result**")
                    txt_download = gr.DownloadButton(label="⬇ TXT", value=None, variant="primary")
                    docx_download = gr.DownloadButton(label="⬇ Word", value=None)
                    json_download = gr.DownloadButton(label="⬇ JSON", value=None)
                    gr.Markdown("TXT for plain text, Word for editing, JSON for structured page results.", elem_id="subtle-note")
            with gr.Accordion("Show page-by-page results", open=False):
                page_results = gr.JSON(label="Page results", open=False)

    upload.change(
        fn=inspect_upload,
        inputs=[upload],
        outputs=[file_info, end_page],
        show_progress="hidden",
    )
    run_button.click(
        fn=run_ocr,
        inputs=[upload, start_page, end_page],
        outputs=[output_text, page_results, txt_download, docx_download, json_download, status],
        show_progress="full",
    )
    check_button.click(fn=check_model, inputs=[], outputs=[status], show_progress="minimal")


if __name__ == "__main__":
    launch_options = {"css": CSS, "theme": gr.themes.Soft(primary_hue="emerald", secondary_hue="teal")}
    if FAVICON_PATH.is_file():
        launch_options["favicon_path"] = str(FAVICON_PATH)
    demo.queue().launch(**launch_options)
